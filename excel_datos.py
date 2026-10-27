"""
Carga y guardado del Excel de alumnos, validación de datos, lista/filtro/
selección en la tabla, y el editor de datos del alumno seleccionado.
"""

import os
from datetime import datetime
import tkinter as tk
from tkinter import filedialog, messagebox
import ttkbootstrap as tb
import pandas as pd
from PIL import Image, ImageOps, ImageTk
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

import estado
from utilidades import (
    sanitizar_nombre_archivo, resolver_ruta_archivo, hacer_ruta_portable,
    actualizar_indicadores, notificar,
)
from respaldos import crear_respaldo


def normalizar_texto_alumno(valor):
    """Limpia espacios y convierte a mayúsculas un dato capturado/importado."""
    if valor is None:
        return ""
    return " ".join(str(valor).strip().split()).upper()


def normalizar_datos_alumnos(df):
    """Normaliza únicamente las columnas académicas y personales del alumno."""
    for columna in estado.COLUMNAS_REQUERIDAS:
        df[columna] = df[columna].map(normalizar_texto_alumno)
    return df


def determinar_ruta_excel_trabajo(ruta):
    """Evita nombres repetidos como *_con_fotos_con_fotos.xlsx."""
    base, _ = os.path.splitext(ruta)
    if base.lower().endswith("_con_fotos"):
        return ruta, True
    return base + "_con_fotos.xlsx", False


def construir_fila_alumno(datos):
    """Normaliza los datos capturados manualmente y agrega campos de control."""
    fila = {
        columna: normalizar_texto_alumno(datos.get(columna, ""))
        for columna in estado.COLUMNAS_REQUERIDAS
    }
    fila["Apellidos"] = " ".join(filter(None, (
        fila["Apellido Paterno"], fila["Apellido Materno"]
    )))
    fila["NombreCompleto"] = " ".join(filter(None, (
        fila["Nombre"], fila["Apellidos"]
    )))
    fila["RutaFoto"] = ""
    fila["FotoTomada"] = "No"
    fila["FechaCaptura"] = ""
    fila["Impresiones"] = 0
    fila["Reposiciones"] = 0
    fila["UltimaImpresion"] = ""
    return fila


# ─── Validación ───────────────────────────────────────────────────────────────
def validar_datos_alumnos(df):
    """Revisa el DataFrame recién cargado en busca de matrículas o CURPs
    vacíos o repetidos. Devuelve una lista de mensajes (vacía si todo bien)."""
    errores = []

    vacias_matricula = int((df["Matricula"] == "").sum())
    if vacias_matricula > 0:
        errores.append(f"• {vacias_matricula} alumno(s) sin Matrícula.")

    vacios_curp = int((df["CURP"] == "").sum())
    if vacios_curp > 0:
        errores.append(f"• {vacios_curp} alumno(s) sin CURP.")

    matriculas_validas = df[df["Matricula"] != ""]["Matricula"]
    repetidas_matricula = sorted(matriculas_validas[matriculas_validas.duplicated(keep=False)].unique())
    if repetidas_matricula:
        errores.append("• Matrículas repetidas: " + ", ".join(repetidas_matricula))

    curps_validos = df[df["CURP"] != ""]["CURP"]
    repetidos_curp = sorted(curps_validos[curps_validos.duplicated(keep=False)].unique())
    if repetidos_curp:
        errores.append("• CURPs repetidos: " + ", ".join(repetidos_curp))

    return errores


# ─── Plantilla descargable ────────────────────────────────────────────────────
def descargar_plantilla_excel():
    """Genera un Excel en blanco con los encabezados requeridos en la primera
    fila, listo para llenar y cargar directamente en CamApp."""
    ruta_destino = filedialog.asksaveasfilename(
        title="Guardar plantilla de Excel",
        defaultextension=".xlsx",
        filetypes=[("Archivos Excel", "*.xlsx")],
        initialfile="Plantilla_Alumnos.xlsx"
    )
    if not ruta_destino:
        return

    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Alumnos"

        relleno_encabezado = PatternFill(start_color="32BEA6", end_color="32BEA6", fill_type="solid")
        for col_idx, nombre_columna in enumerate(estado.COLUMNAS_REQUERIDAS, start=1):
            celda = ws.cell(row=1, column=col_idx, value=nombre_columna)
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = relleno_encabezado
            celda.alignment = Alignment(horizontal="center")

        for col_idx, nombre_columna in enumerate(estado.COLUMNAS_REQUERIDAS, start=1):
            ancho = len(nombre_columna) + 4
            ws.column_dimensions[get_column_letter(col_idx)].width = ancho

        ws.freeze_panes = "A2"
        wb.save(ruta_destino)

        messagebox.showinfo(
            "Plantilla generada",
            f"Plantilla guardada en:\n{ruta_destino}\n\n"
            f"Compártela con la escuela para que la llenen desde la fila 2, "
            f"respetando el nombre y orden de las columnas."
        )
    except PermissionError:
        messagebox.showerror(
            "Error al guardar",
            "No se pudo guardar la plantilla. Verifica que no haya un archivo "
            "abierto con ese mismo nombre en Excel y vuelve a intentar."
        )
    except Exception as e:
        messagebox.showerror("Error al generar la plantilla", str(e))


# ─── Carga y guardado del Excel de trabajo ───────────────────────────────────
def cargar_excel(ruta_forzada=None, preguntar_progreso=True):
    ruta = ruta_forzada or filedialog.askopenfilename(
        title="Selecciona el Excel de alumnos",
        filetypes=[("Archivos Excel", "*.xlsx *.xls")]
    )
    if not ruta:
        return

    if estado.ruta_proyecto:
        ruta_trabajo_posible = estado.ruta_excel_trabajo
        es_archivo_trabajo = (
            os.path.normcase(os.path.abspath(ruta))
            == os.path.normcase(os.path.abspath(ruta_trabajo_posible))
        )
    else:
        ruta_trabajo_posible, es_archivo_trabajo = determinar_ruta_excel_trabajo(ruta)

    usar_progreso_previo = False
    if (preguntar_progreso and not es_archivo_trabajo
            and os.path.exists(ruta_trabajo_posible)):
        usar_progreso_previo = messagebox.askyesno(
            "Progreso encontrado",
            f"Se encontró un archivo de trabajo previo:\n"
            f"{os.path.basename(ruta_trabajo_posible)}\n\n"
            f"¿Deseas continuar desde ahí (conservando las fotos ya tomadas)?\n\n"
            f"Si eliges 'No', se volverá a partir del Excel original que seleccionaste."
        )

    ruta_leer = (
        ruta
        if es_archivo_trabajo
        else ruta_trabajo_posible if usar_progreso_previo else ruta
    )

    try:
        df = pd.read_excel(ruta_leer, dtype=str)
    except Exception as e:
        messagebox.showerror("Error al leer el Excel", str(e))
        return

    df.columns = [str(c).strip() for c in df.columns]

    faltantes = [c for c in estado.COLUMNAS_REQUERIDAS if c not in df.columns]
    if faltantes:
        messagebox.showerror(
            "Columnas faltantes",
            "El Excel no tiene las siguientes columnas requeridas:\n\n"
            + ", ".join(faltantes) +
            "\n\nRevisa los encabezados y vuelve a intentar."
        )
        return

    df = df.fillna("")
    normalizar_datos_alumnos(df)

    errores_validacion = validar_datos_alumnos(df)
    if errores_validacion:
        messagebox.showerror(
            "Datos inválidos en el Excel",
            "Se encontraron los siguientes problemas antes de cargar:\n\n"
            + "\n".join(errores_validacion) +
            "\n\nCorrige el Excel (matrículas y CURPs deben ser únicos y no estar "
            "vacíos) y vuelve a intentar."
        )
        return

    # Siempre se recalculan para actualizar también archivos de trabajo antiguos.
    df["Apellidos"] = (
        df["Apellido Paterno"] + " " + df["Apellido Materno"]
    ).str.replace(r"\s+", " ", regex=True).str.strip()
    df["NombreCompleto"] = (
        df["Nombre"] + " " + df["Apellidos"]
    ).str.replace(r"\s+", " ", regex=True).str.strip()

    if "RutaFoto" not in df.columns:
        df["RutaFoto"] = ""
    if "FotoTomada" not in df.columns:
        df["FotoTomada"] = "No"
    if "FechaCaptura" not in df.columns:
        df["FechaCaptura"] = ""
    if "Impresiones" not in df.columns:
        df["Impresiones"] = 0
    if "Reposiciones" not in df.columns:
        df["Reposiciones"] = 0
    if "UltimaImpresion" not in df.columns:
        df["UltimaImpresion"] = ""
    for columna in ("Impresiones", "Reposiciones"):
        df[columna] = pd.to_numeric(df[columna], errors="coerce").fillna(0).astype(int)

    estado.df_alumnos = df
    estado.ruta_excel_trabajo = ruta_trabajo_posible
    estado.alumno_actual_idx = None
    if not estado.ruta_proyecto:
        estado.FOTOS_DIR = os.path.join(os.path.dirname(ruta), "fotos_alumnos")
    os.makedirs(estado.FOTOS_DIR, exist_ok=True)

    grupos = sorted([g for g in estado.df_alumnos["Grupo"].unique().tolist() if g])
    estado.combo_grupo["values"] = ["Todos"] + grupos
    estado.combo_grupo.set("Todos")
    actualizar_opciones_filtros()
    if estado.combo_estado_foto:
        estado.combo_estado_foto.set("Todos")
    estado.entry_buscar.delete(0, tk.END)

    # Crear una subcarpeta por cada grupo dentro de fotos_alumnos/
    for grupo in grupos:
        carpeta_grupo = os.path.join(estado.FOTOS_DIR, sanitizar_nombre_archivo(grupo))
        os.makedirs(carpeta_grupo, exist_ok=True)

    guardar_excel(mostrar_aviso=False)
    actualizar_lista()
    actualizar_etiqueta_alumno()


def actualizar_etiqueta_excel():
    if estado.df_alumnos is None:
        estado.etiqueta_excel.configure(text="Ningún archivo cargado")
        return
    total = len(estado.df_alumnos)
    con_foto = (estado.df_alumnos["FotoTomada"] == "Sí").sum()
    estado.etiqueta_excel.configure(
        text=f"{total} alumnos cargados  |  {con_foto} con foto  |  "
             f"Guardando en: {os.path.basename(estado.ruta_excel_trabajo)}"
    )


def guardar_excel(mostrar_aviso=True):
    if estado.df_alumnos is None or estado.ruta_excel_trabajo is None:
        if mostrar_aviso:
            messagebox.showwarning("Sin datos", "Primero carga un Excel de alumnos.")
        return
    try:
        crear_respaldo(estado.ruta_excel_trabajo, "excel")
        estado.df_alumnos.to_excel(estado.ruta_excel_trabajo, index=False)
        actualizar_etiqueta_excel()
        if mostrar_aviso:
            messagebox.showinfo("Guardado", f"Excel actualizado:\n{estado.ruta_excel_trabajo}")
    except PermissionError:
        messagebox.showerror(
            "Error al guardar",
            "No se pudo guardar el Excel. Verifica que no esté abierto en otro "
            "programa (por ejemplo Excel) y vuelve a intentar."
        )
    except Exception as e:
        messagebox.showerror("Error al guardar", str(e))


# ─── Lista / selección de alumnos ────────────────────────────────────────────
def actualizar_lista_debounced(*_):
    """Espera una pequeña pausa después de la última tecla antes de refrescar
    la tabla, para no reconstruirla completa en cada pulsación (importante
    cuando hay cientos de alumnos cargados con el filtro en 'Todos')."""
    if estado.id_busqueda_pendiente is not None:
        estado.root.after_cancel(estado.id_busqueda_pendiente)
    estado.id_busqueda_pendiente = estado.root.after(250, actualizar_lista)


def actualizar_opciones_filtros():
    """Actualiza las listas desplegables que dependen del Excel cargado."""
    if estado.df_alumnos is None:
        return
    for columna in ("especialidad", "grupo", "turno", "generacion", "semestre"):
        combo = estado.filtros_tabla.get(columna)
        nombre = {
            "especialidad": "Especialidad", "grupo": "Grupo", "turno": "Turno",
            "generacion": "Generacion", "semestre": "Semestre",
        }[columna]
        if combo is not None and nombre in estado.df_alumnos.columns:
            valores = sorted({str(valor) for valor in estado.df_alumnos[nombre].dropna() if str(valor).strip()})
            combo.configure(values=["Todos"] + valores)
            if not combo.get():
                combo.set("Todos")


def actualizar_lista(*_):
    for item in estado.tree.get_children():
        estado.tree.delete(item)

    if estado.df_alumnos is None:
        actualizar_indicadores()
        return

    grupo_sel = estado.combo_grupo.get()
    estado_foto_sel = estado.combo_estado_foto.get() if estado.combo_estado_foto else "Todos"
    texto_busqueda = estado.entry_buscar.get().strip().lower()

    df_filtrado = estado.df_alumnos
    if grupo_sel and grupo_sel != "Todos":
        df_filtrado = df_filtrado[df_filtrado["Grupo"] == grupo_sel]
    if estado_foto_sel == "Con foto":
        df_filtrado = df_filtrado[df_filtrado["FotoTomada"] == "Sí"]
    elif estado_foto_sel == "Sin foto":
        df_filtrado = df_filtrado[df_filtrado["FotoTomada"] != "Sí"]

    if texto_busqueda:
        mask = (
            df_filtrado["NombreCompleto"].str.lower().str.contains(texto_busqueda, na=False) |
            df_filtrado["Matricula"].str.lower().str.contains(texto_busqueda, na=False)
        )
        df_filtrado = df_filtrado[mask]

    filtros = estado.filtros_tabla
    for clave, nombre in {
        "nombre": "Nombre", "apellido_paterno": "Apellido Paterno",
        "apellido_materno": "Apellido Materno", "matricula": "Matricula",
    }.items():
        campo = filtros.get(clave)
        texto = campo.get().strip().lower() if campo is not None else ""
        if texto:
            df_filtrado = df_filtrado[
                df_filtrado[nombre].astype(str).str.lower().str.contains(texto, na=False)]
    for clave, nombre in {
        "especialidad": "Especialidad", "grupo": "Grupo", "turno": "Turno",
        "generacion": "Generacion", "semestre": "Semestre",
    }.items():
        combo = filtros.get(clave)
        valor = combo.get() if combo is not None else "Todos"
        if valor and valor != "Todos":
            df_filtrado = df_filtrado[df_filtrado[nombre].astype(str) == valor]
    filtro_foto = filtros.get("foto")
    valor_foto = filtro_foto.get() if filtro_foto is not None else "Todos"
    if valor_foto == "Con foto":
        df_filtrado = df_filtrado[df_filtrado["FotoTomada"] == "Sí"]
    elif valor_foto == "Sin foto":
        df_filtrado = df_filtrado[df_filtrado["FotoTomada"] != "Sí"]

    for idx, fila in df_filtrado.iterrows():
        tiene_foto = fila["FotoTomada"] == "Sí"
        texto_estado = "✔ Sí" if tiene_foto else "— No"
        tag = "con_foto" if tiene_foto else "sin_foto"
        estado.tree.insert("", "end", iid=str(idx), values=(
            fila["Nombre"], fila["Apellido Paterno"], fila["Apellido Materno"],
            fila["Matricula"], fila["Especialidad"],
            fila["Grupo"], fila["Turno"], fila["Generacion"], fila["Semestre"],
            texto_estado,
        ), tags=(tag,))
    actualizar_indicadores()


def seleccionar_alumno(event=None):
    seleccion = estado.tree.selection()
    if not seleccion:
        estado.alumno_actual_idx = None
        if estado.boton_imprimir_credencial:
            estado.boton_imprimir_credencial.configure(state="disabled")
        for boton in (estado.boton_cargar_foto, estado.boton_historial):
            if boton:
                boton.configure(state="disabled")
        actualizar_etiqueta_alumno()
        return
    estado.alumno_actual_idx = int(seleccion[0])
    if estado.boton_imprimir_credencial:
        estado.boton_imprimir_credencial.configure(state="normal")
    for boton in (estado.boton_cargar_foto, estado.boton_historial):
        if boton:
            boton.configure(state="normal")
    actualizar_etiqueta_alumno()


def seleccionar_siguiente_pendiente(idx_recien_guardado):
    """Tras guardar una foto, selecciona automáticamente al siguiente alumno
    pendiente dentro de la vista filtrada actual (mismo grupo/búsqueda activos).
    Si no queda ninguno pendiente en esa vista, deja la lista sin selección
    para que el operador elija manualmente."""
    hijos = estado.tree.get_children()
    if not hijos:
        return

    ids_hijos = [int(iid) for iid in hijos]

    try:
        posicion_actual = ids_hijos.index(idx_recien_guardado)
    except ValueError:
        posicion_actual = -1  # el alumno recién guardado ya no aparece en la vista actual

    orden_busqueda = ids_hijos[posicion_actual + 1:] + ids_hijos[:posicion_actual + 1]

    for idx_candidato in orden_busqueda:
        if estado.df_alumnos.at[idx_candidato, "FotoTomada"] != "Sí":
            iid_candidato = str(idx_candidato)
            estado.tree.selection_set(iid_candidato)
            estado.tree.see(iid_candidato)
            return

    seleccion_actual = estado.tree.selection()
    if seleccion_actual:
        estado.tree.selection_remove(seleccion_actual)
    else:
        seleccionar_alumno()  # fuerza limpiar alumno_actual_idx aunque ya no había selección


def actualizar_etiqueta_alumno():
    if estado.alumno_actual_idx is None or estado.df_alumnos is None:
        estado.etiqueta_alumno.configure(text="Sin alumno seleccionado")
        actualizar_previsualizacion_foto()
        return
    fila = estado.df_alumnos.loc[estado.alumno_actual_idx]
    estado.etiqueta_alumno.configure(
        text=f"Alumno seleccionado: {fila['NombreCompleto']}  |  "
             f"Matrícula: {fila['Matricula']}  |  Grupo: {fila['Grupo']}  |  "
             f"Impresiones: {int(fila.get('Impresiones', 0) or 0)}  |  "
             f"Reposiciones: {int(fila.get('Reposiciones', 0) or 0)}"
    )
    actualizar_previsualizacion_foto()


def actualizar_previsualizacion_foto():
    """Muestra la fotografía del alumno seleccionado, si existe."""
    etiqueta = estado.etiqueta_foto
    if etiqueta is None:
        return
    etiqueta.configure(image="", text="Sin fotografía")
    etiqueta.image = None
    if estado.alumno_actual_idx is None or estado.df_alumnos is None:
        return
    fila = estado.df_alumnos.loc[estado.alumno_actual_idx]
    ruta = resolver_ruta_archivo(fila.get("RutaFoto", ""))
    if not ruta or not os.path.isfile(ruta):
        return
    try:
        with Image.open(ruta) as imagen:
            imagen = ImageOps.contain(imagen.convert("RGB"), (250, 300))
            lienzo = Image.new("RGB", (260, 320), "white")
            posicion = (
                (lienzo.width - imagen.width) // 2,
                (lienzo.height - imagen.height) // 2,
            )
            lienzo.paste(imagen, posicion)
        imagen_tk = ImageTk.PhotoImage(lienzo)
        etiqueta.configure(image=imagen_tk, text="")
        etiqueta.image = imagen_tk
    except (OSError, ValueError):
        etiqueta.configure(text="No se pudo cargar la fotografía")


# ─── Alta manual de alumnos ──────────────────────────────────────────────────
def abrir_agregar_alumno():
    """Agrega un alumno al Excel activo o crea una base de trabajo nueva."""
    ventana = tk.Toplevel(estado.root)
    ventana.title("Agregar Alumno")
    ventana.resizable(False, False)
    ventana.transient(estado.root)
    ventana.grab_set()

    tb.Label(
        ventana,
        text="Agregar alumno manualmente",
        font=("Segoe UI", 12, "bold"),
        bootstyle="primary",
    ).grid(row=0, column=0, columnspan=2, padx=20, pady=(18, 5))

    texto_destino = (
        f"Se agregará a: {os.path.basename(estado.ruta_excel_trabajo)}"
        if estado.df_alumnos is not None and estado.ruta_excel_trabajo
        else "Se creará un nuevo Excel de trabajo para este alumno."
    )
    tb.Label(
        ventana, text=texto_destino, bootstyle="secondary",
        font=("Segoe UI", 9, "italic")
    ).grid(row=1, column=0, columnspan=2, padx=20, pady=(0, 12))

    campos = [
        "Matricula", "Nombre", "Apellido Paterno", "Apellido Materno",
        "CURP", "Grupo", "Turno", "Especialidad", "Semestre", "Generacion",
    ]
    valores_iniciales = {
        "Turno": "MATUTINO",
        "Especialidad": "N/A",
    }
    entradas = {}

    for numero_fila, campo in enumerate(campos, start=2):
        tb.Label(ventana, text=f"{campo}:").grid(
            row=numero_fila, column=0, sticky="e", padx=(20, 6), pady=4)
        entrada = tb.Entry(ventana, width=34)
        entrada.insert(0, valores_iniciales.get(campo, ""))
        entrada.grid(row=numero_fila, column=1, padx=(0, 20), pady=4)
        entradas[campo] = entrada

    entradas["Matricula"].focus_set()

    def guardar_alumno():
        datos = {campo: entrada.get() for campo, entrada in entradas.items()}
        fila_nueva = construir_fila_alumno(datos)

        obligatorios = ["Matricula", "Nombre", "Apellido Paterno", "CURP", "Grupo"]
        faltantes = [campo for campo in obligatorios if not fila_nueva[campo]]
        if faltantes:
            messagebox.showwarning(
                "Datos incompletos",
                "Completa los siguientes campos:\n\n" + ", ".join(faltantes),
                parent=ventana,
            )
            return

        if estado.df_alumnos is not None:
            if (estado.df_alumnos["Matricula"] == fila_nueva["Matricula"]).any():
                messagebox.showwarning(
                    "Matrícula duplicada",
                    "Ya existe un alumno con esa matrícula.",
                    parent=ventana,
                )
                return
            if (estado.df_alumnos["CURP"] == fila_nueva["CURP"]).any():
                messagebox.showwarning(
                    "CURP duplicado",
                    "Ya existe un alumno con ese CURP.",
                    parent=ventana,
                )
                return

        if estado.df_alumnos is None:
            ruta_nueva = (
                estado.ruta_excel_trabajo
                if estado.ruta_proyecto
                else filedialog.asksaveasfilename(
                    title="Guardar nuevo Excel de trabajo",
                    defaultextension=".xlsx",
                    filetypes=[("Archivos Excel", "*.xlsx")],
                    initialfile="Alumnos_Manuales_con_fotos.xlsx",
                    parent=ventana,
                )
            )
            if not ruta_nueva:
                return
            estado.ruta_excel_trabajo = ruta_nueva
            if not estado.ruta_proyecto:
                estado.FOTOS_DIR = os.path.join(
                    os.path.dirname(ruta_nueva), "fotos_alumnos")
            os.makedirs(estado.FOTOS_DIR, exist_ok=True)
            columnas = estado.COLUMNAS_REQUERIDAS + [
                "Apellidos", "NombreCompleto", "RutaFoto",
                "FotoTomada", "FechaCaptura", "Impresiones",
                "Reposiciones", "UltimaImpresion",
            ]
            estado.df_alumnos = pd.DataFrame(columns=columnas)

        estado.df_alumnos = pd.concat(
            [estado.df_alumnos, pd.DataFrame([fila_nueva])],
            ignore_index=True,
        ).fillna("")
        nuevo_idx = len(estado.df_alumnos) - 1

        grupos = sorted(
            g for g in estado.df_alumnos["Grupo"].unique().tolist() if g)
        estado.combo_grupo["values"] = ["Todos"] + grupos
        estado.combo_grupo.set("Todos")
        actualizar_opciones_filtros()
        estado.entry_buscar.delete(0, tk.END)

        carpeta_grupo = os.path.join(
            estado.FOTOS_DIR, sanitizar_nombre_archivo(fila_nueva["Grupo"]))
        os.makedirs(carpeta_grupo, exist_ok=True)

        guardar_excel(mostrar_aviso=False)
        actualizar_lista()
        estado.tree.selection_set(str(nuevo_idx))
        estado.tree.see(str(nuevo_idx))
        seleccionar_alumno()

        ventana.destroy()
        notificar(
            f"{fila_nueva['NombreCompleto']} fue agregado correctamente",
            "success")

    frame_botones = tb.Frame(ventana)
    frame_botones.grid(
        row=len(campos) + 2, column=0, columnspan=2, pady=(14, 18))
    tb.Button(
        frame_botones, text="Agregar Alumno", bootstyle="success",
        cursor="hand2", command=guardar_alumno
    ).pack(side="left", padx=6, ipady=4)
    tb.Button(
        frame_botones, text="Cancelar", bootstyle="secondary",
        cursor="hand2", command=ventana.destroy
    ).pack(side="left", padx=6, ipady=4)

    ventana.bind("<Return>", lambda event: guardar_alumno())
    ventana.bind("<Escape>", lambda event: ventana.destroy())


# ─── Editor de datos del alumno ──────────────────────────────────────────────
def abrir_editor_alumno():
    """Abre un formulario para corregir Nombre, Apellidos y CURP del alumno
    seleccionado (la Matrícula no es editable: es la clave que usa todo el
    programa para identificar al alumno y sus archivos)."""
    if estado.alumno_actual_idx is None or estado.df_alumnos is None:
        messagebox.showwarning("Sin alumno", "Selecciona un alumno de la lista primero.")
        return

    fila = estado.df_alumnos.loc[estado.alumno_actual_idx]
    campos = ["Nombre", "Apellido Paterno", "Apellido Materno", "CURP"]

    ventana_editar = tk.Toplevel(estado.root)
    ventana_editar.title("Editar Datos del Alumno")
    ventana_editar.resizable(False, False)

    tb.Label(ventana_editar, text=f"Matrícula: {fila['Matricula']}  (no editable)",
             font=("Segoe UI", 9, "italic"), bootstyle="secondary"
             ).grid(row=0, column=0, columnspan=2, padx=15, pady=(15, 10), sticky="w")

    entradas = {}
    for i, campo in enumerate(campos, start=1):
        tb.Label(ventana_editar, text=f"{campo}:").grid(row=i, column=0, sticky="e", padx=(15, 5), pady=6)
        entrada = tb.Entry(ventana_editar, width=30)
        entrada.insert(0, fila[campo])
        entrada.grid(row=i, column=1, padx=(0, 15), pady=6)
        entradas[campo] = entrada
    entradas["Nombre"].focus_set()

    def guardar_cambios():
        nuevos_valores = {
            campo: normalizar_texto_alumno(entradas[campo].get())
            for campo in campos
        }

        if not nuevos_valores["Nombre"] or not nuevos_valores["Apellido Paterno"]:
            messagebox.showwarning("Datos incompletos",
                "Nombre y Apellido Paterno no pueden quedar vacíos.",
                parent=ventana_editar)
            return

        if not nuevos_valores["CURP"]:
            messagebox.showwarning("Datos incompletos",
                "El CURP no puede quedar vacío.", parent=ventana_editar)
            return

        otros = estado.df_alumnos.drop(index=estado.alumno_actual_idx)
        if (otros["CURP"] == nuevos_valores["CURP"]).any():
            messagebox.showwarning("CURP duplicado",
                "Ese CURP ya pertenece a otro alumno en la lista.", parent=ventana_editar)
            return

        nombre_completo_anterior = fila["NombreCompleto"]
        ruta_foto_guardada = fila["RutaFoto"]
        ruta_foto_anterior = resolver_ruta_archivo(ruta_foto_guardada)

        for campo, valor in nuevos_valores.items():
            estado.df_alumnos.at[estado.alumno_actual_idx, campo] = valor

        nuevo_nombre_completo = " ".join((
            nuevos_valores["Nombre"] + " " +
            nuevos_valores["Apellido Paterno"] + " " +
            nuevos_valores["Apellido Materno"]
        ).split())
        nuevos_apellidos = " ".join((
            nuevos_valores["Apellido Paterno"] + " " +
            nuevos_valores["Apellido Materno"]
        ).split())
        estado.df_alumnos.at[estado.alumno_actual_idx, "Apellidos"] = nuevos_apellidos
        estado.df_alumnos.at[estado.alumno_actual_idx, "NombreCompleto"] = nuevo_nombre_completo

        # Si ya tenía foto y el nombre cambió, renombrar el archivo físico
        # para que siga coincidiendo con el nombre corregido.
        if (ruta_foto_anterior and os.path.exists(ruta_foto_anterior)
                and nuevo_nombre_completo != nombre_completo_anterior):
            carpeta = os.path.dirname(ruta_foto_anterior)
            nuevo_nombre_archivo = sanitizar_nombre_archivo(
                f"{nuevo_nombre_completo}_{fila['Matricula']}") + ".png"
            nueva_ruta = os.path.join(carpeta, nuevo_nombre_archivo)
            try:
                os.rename(ruta_foto_anterior, nueva_ruta)
                estado.df_alumnos.at[estado.alumno_actual_idx, "RutaFoto"] = (
                    hacer_ruta_portable(nueva_ruta))
            except OSError:
                pass  # si falla el renombrado, la ruta anterior sigue siendo válida

        guardar_excel(mostrar_aviso=False)
        actualizar_lista()
        actualizar_etiqueta_alumno()
        ventana_editar.destroy()
        notificar("Datos del alumno actualizados", "success")

    frame_btns_editar = tb.Frame(ventana_editar)
    frame_btns_editar.grid(row=len(campos) + 1, column=0, columnspan=2, pady=15)

    tb.Button(frame_btns_editar, text="Guardar Cambios", bootstyle="success",
              cursor="hand2", command=guardar_cambios).grid(row=0, column=0, padx=8, ipady=4)
    tb.Button(frame_btns_editar, text="Cancelar", bootstyle="secondary",
              cursor="hand2", command=ventana_editar.destroy).grid(row=0, column=1, padx=8, ipady=4)

    ventana_editar.bind('<Return>', lambda e: guardar_cambios())
    ventana_editar.bind('<Escape>', lambda e: ventana_editar.destroy())
